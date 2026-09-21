"""Builds the single 'Overview Workflow' Excel deliverable (.xlsx).

Per current requirements this is the ONLY document output the platform
produces (DOCX/PDF reports have been removed). Everything here is derived
purely from the parsed XML/JSON domain model (RepositoryModel) via the
deterministic analyzer functions -- no data is invented.

Workbook layout:
  Tab 1 "Session-Mapping"     : Session, Mapping_Name, Mapplet, prev_session,
                                 next_Session, prev_mapping, next_Mapping
  Tab 2 "Table-Transformation": Mapping Name, Mapplet, Source Table, Target Table,
                                 Transformation lineage, Unused Transformation
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.analyzer.workflow_analyzer import (
    compute_execution_order, build_mapping_lineage_by_source_qualifier, find_unused_transformations,
)

NAVY = "1F3864"
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
HEADER_FILL = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
BODY_FONT = Font(size=10)
WRAP_TOP = Alignment(vertical="top", wrap_text=True)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
STRIPE_FILL = PatternFill(start_color="F2F6FA", end_color="F2F6FA", fill_type="solid")

NONE_LABEL = "(none)"


def _write_sheet(ws, headers, rows, col_widths):
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)
        cell.border = BORDER
    for r_idx, row in enumerate(rows, start=2):
        for c_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            cell.font = BODY_FONT
            cell.alignment = WRAP_TOP
            cell.border = BORDER
            if r_idx % 2 == 0:
                cell.fill = STRIPE_FILL
    for c_idx, width in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(c_idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(len(rows) + 1, 1)}"


def _mapplets_for_mapping_str(repo, mapping_name):
    """Comma-joined mapplet names owned by a mapping, or NONE_LABEL if the
    mapping uses none / isn't found."""
    mapping = repo.mappings.get(mapping_name) if mapping_name else None
    if mapping is None or not mapping.mapplets:
        return NONE_LABEL
    return ", ".join(mapping.mapplets)


def _session_mapping_rows(repo):
    wf = repo.workflow
    order = wf.execution_order or compute_execution_order(repo)
    session_order = [t for t in order if t in wf.sessions]

    rows = []
    for idx, name in enumerate(session_order):
        sess = wf.sessions[name]
        prev_session = session_order[idx - 1] if idx > 0 else NONE_LABEL
        next_session = session_order[idx + 1] if idx < len(session_order) - 1 else NONE_LABEL
        prev_mapping = wf.sessions[prev_session].mapping_name if prev_session != NONE_LABEL else NONE_LABEL
        next_mapping = wf.sessions[next_session].mapping_name if next_session != NONE_LABEL else NONE_LABEL
        rows.append([
            sess.name,
            sess.mapping_name or NONE_LABEL,
            _mapplets_for_mapping_str(repo, sess.mapping_name),
            prev_session,
            next_session,
            prev_mapping,
            next_mapping,
        ])
    return rows


def _mapping_execution_order(repo):
    """Mapping names ordered by the order their owning sessions execute
    (first occurrence wins); any mapping with no session is appended,
    sorted, at the end so nothing is dropped."""
    wf = repo.workflow
    order = wf.execution_order or compute_execution_order(repo)
    session_order = [t for t in order if t in wf.sessions]

    seen = set()
    mapping_order = []
    for name in session_order:
        m_name = wf.sessions[name].mapping_name
        if m_name and m_name not in seen:
            seen.add(m_name)
            mapping_order.append(m_name)
    for m_name in sorted(repo.mappings.keys()):
        if m_name not in seen:
            mapping_order.append(m_name)
    return mapping_order


def _table_transformation_rows(repo):
    """One row per Source Qualifier (not per mapping): a mapping can have
    n sources feeding one SQ and multiple independent SQ branches reaching
    different targets, so each SQ branch gets its own row with just its own
    source(s), the target(s) it actually reaches, and its own full lineage.

    "Unused Transformation" is a mapping-level fact (a transformation with
    zero connector edges can't belong to any one SQ branch), so it's computed
    once per mapping and shown only on that mapping's first row; the other
    rows for the same mapping leave it blank rather than repeating it."""
    rows = []
    for m_name in _mapping_execution_order(repo):
        mapping = repo.mappings.get(m_name)
        mapplet_str = _mapplets_for_mapping_str(repo, m_name)
        if mapping is None:
            rows.append([m_name, mapplet_str, NONE_LABEL, NONE_LABEL, "(mapping not found)", NONE_LABEL])
            continue
        unused = find_unused_transformations(repo, m_name)
        unused_str = ", ".join(unused) if unused else NONE_LABEL
        sq_branches = build_mapping_lineage_by_source_qualifier(repo, m_name)
        if not sq_branches:
            rows.append([m_name, mapplet_str, NONE_LABEL, NONE_LABEL, NONE_LABEL, unused_str])
            continue
        for idx, branch in enumerate(sq_branches):
            source_table = ", ".join(branch["sources"]) or NONE_LABEL
            target_table = ", ".join(branch["targets"]) or NONE_LABEL
            lineage_str = " -> ".join(branch["lineage"]) if branch["lineage"] else NONE_LABEL
            rows.append([m_name, mapplet_str if idx == 0 else "", source_table, target_table, lineage_str,
                         unused_str if idx == 0 else ""])
    return rows


def generate_overview_workflow_excel(repo, out_path: str) -> str:
    wb = Workbook()

    ws1 = wb.active
    ws1.title = "Session-Mapping"
    _write_sheet(
        ws1,
        ["Session", "Mapping_Name", "Mapplet", "prev_session", "next_Session", "prev_mapping", "next_Mapping"],
        _session_mapping_rows(repo),
        col_widths=[28, 28, 26, 28, 28, 28, 28],
    )

    ws2 = wb.create_sheet("Table-Transformation")
    _write_sheet(
        ws2,
        ["Mapping Name", "Mapplet", "Source Table", "Target Table", "Transformation lineage", "Unused Transformation"],
        _table_transformation_rows(repo),
        col_widths=[28, 26, 26, 26, 90, 32],
    )

    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    wb.save(out_path)
    return out_path
