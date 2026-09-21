"""Generate the multi-sheet workbook used by the Informatica XML Downloads page.

Workbook sheets:
1) Field_Lineage
2) Referrential-other Tables
3) Eligibility Rules
4) Eligibility Rules - Summary
"""
from __future__ import annotations

import os
import re
from collections import defaultdict
from typing import Dict, Iterable, List, Optional, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.analyzer import field_lineage_analyzer as fla
from app.analyzer.workflow_analyzer import compute_execution_order

NAVY = "1F3864"
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
HEADER_FILL = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
BODY_FONT = Font(size=10)
WRAP_TOP = Alignment(vertical="top", wrap_text=True)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
STRIPE_FILL = PatternFill(start_color="F2F6FA", end_color="F2F6FA", fill_type="solid")


def _write_sheet(ws, headers: List[str], rows: Iterable[Dict[str, object]], col_widths: List[int]) -> None:
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)
        cell.border = BORDER

    row_count = 1
    for r_idx, row in enumerate(rows, start=2):
        row_count = r_idx
        for c_idx, col_name in enumerate(headers, start=1):
            value = row.get(col_name, "")
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            cell.font = BODY_FONT
            cell.alignment = WRAP_TOP
            cell.border = BORDER
            if r_idx % 2 == 0:
                cell.fill = STRIPE_FILL

    for c_idx, width in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(c_idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(row_count, 1)}"


def _ordered_mapping_sessions(repo) -> List[Tuple[str, Optional[str]]]:
    wf = repo.workflow
    ordered: List[Tuple[str, Optional[str]]] = []
    seen = set()

    if wf is not None:
        order = wf.execution_order or compute_execution_order(repo)
        session_order = [task for task in order if task in wf.sessions] or list(wf.sessions.keys())
        for sname in session_order:
            mname = wf.sessions[sname].mapping_name
            if mname in repo.mappings and mname not in seen:
                ordered.append((mname, sname))
                seen.add(mname)

    for mname in repo.mappings:
        if mname not in seen:
            ordered.append((mname, fla.session_for_mapping(repo, mname)))
            seen.add(mname)
    return ordered


def _field_lineage_rows(repo) -> List[Dict[str, object]]:
    flat = fla.build_full_workflow_field_lineage(repo)
    rows = []
    for r in flat:
        rows.append({
            "Mapping": r.get("Mapping", ""),
            "Session": r.get("Session", ""),
            "Source Table": r.get("Source Table", ""),
            "Source Field": r.get("Source Field", ""),
            "Target Table": r.get("Target Table", ""),
            "Target Field": r.get("Target Field", ""),
            "Transformation": r.get("Transformation", ""),
            "Transformation Type": r.get("Transformation Type", ""),
            "Transformation_Full_Lineage_path": r.get("Transformation_Full_Lineage_Path", ""),
            "Individual_Transformations": (
                r.get("Individual_Transformations")
                or r.get("Individual_Transformation")
                or ""
            ),
            "Links": r.get("Links", ""),
            "Hop count": r.get("Hop Count", ""),
        })
    return rows


_TOKEN_RE = re.compile(r"\b[A-Za-z_][A-Za-z0-9_$#]*(?:\.[A-Za-z_][A-Za-z0-9_$#]*)?\b")
_OPER_RE = re.compile(r"(<=|>=|<>|!=|=|<|>)")
_SQL_SELECT_RE = re.compile(r"\bselect\b(.*?)\bfrom\b", re.IGNORECASE | re.DOTALL)
_KEYWORDS = {
    "AND", "OR", "NOT", "IN", "IS", "NULL", "LIKE", "BETWEEN", "CASE",
    "WHEN", "THEN", "ELSE", "END", "UPPER", "LOWER", "TRIM", "LTRIM", "RTRIM",
    "IIF", "DECODE", "NVL", "TO_DATE", "TO_CHAR", "ABS", "ROUND", "SUBSTR",
    "LENGTH", "REPLACE", "REPLACECHR",
}


def _normalize_field_token(token: str) -> str:
    if "." in token:
        return token.split(".")[-1]
    return token


def _extract_fields_from_condition(text: str) -> List[str]:
    if not text:
        return []
    fields = []
    for part in re.split(r"\bAND\b|\bOR\b", text, flags=re.IGNORECASE):
        if not _OPER_RE.search(part):
            continue
        for token in _TOKEN_RE.findall(part):
            t = token.strip()
            if t.upper() in _KEYWORDS:
                continue
            if re.fullmatch(r"[0-9]+", t):
                continue
            fields.append(_normalize_field_token(t))
    return fields


def _extract_fields_from_sql_override(sql: str) -> List[str]:
    if not sql:
        return []
    m = _SQL_SELECT_RE.search(sql)
    if not m:
        return []
    selected = m.group(1)
    fields = []
    for col_expr in selected.split(","):
        expr = col_expr.strip()
        if not expr:
            continue
        # Remove common alias forms.
        expr = re.sub(r"\s+AS\s+\w+$", "", expr, flags=re.IGNORECASE)
        expr = re.sub(r"\s+\w+$", "", expr)
        tokens = _TOKEN_RE.findall(expr)
        if tokens:
            fields.append(_normalize_field_token(tokens[-1]))
    return fields


def _reference_other_tables_rows(repo) -> List[Dict[str, object]]:
    rows: List[Dict[str, object]] = []
    seen = set()

    for mapping_name, _session_name in _ordered_mapping_sessions(repo):
        mapping = repo.mappings.get(mapping_name)
        if mapping is None:
            continue
        for inst in mapping.instances:
            if inst.type != "TRANSFORMATION":
                continue
            tobj = repo.transformation(mapping_name, inst.ref_name)
            if tobj is None:
                continue
            if "lookup" not in (tobj.type or "").lower():
                continue

            attrs = tobj.attributes or {}
            table_name = attrs.get("Lookup table name") or attrs.get("Lookup source name") or ""
            alias = attrs.get("Lookup source name") if attrs.get("Lookup source name") != table_name else ""

            fields = set()
            fields.update(_extract_fields_from_condition(attrs.get("Lookup Condition", "")))
            fields.update(_extract_fields_from_condition(attrs.get("Lookup source filter", "")))
            fields.update(_extract_fields_from_sql_override(attrs.get("Lookup SQL Override", "")))

            if not fields:
                fields = {""}

            for field_name in sorted(fields, key=lambda x: (x is None, str(x))):
                key = (mapping_name, inst.name, table_name, alias, field_name)
                if key in seen:
                    continue
                seen.add(key)
                rows.append({
                    "Mapping": mapping_name,
                    "Transformation": inst.name,
                    "Referrential Table Name": table_name,
                    "Alias /Original Name": alias,
                    "Referrential Field": field_name,
                })
    return rows


def _is_rule_expression(expr: str) -> bool:
    text = (expr or "").strip().lower()
    if not text:
        return False
    return any(k in text for k in ("iif(", "decode(", " case ", " and ", " or "))


def _plain_language_for_expression(port_name: str) -> str:
    return f"Conditionally derive {port_name} according to the expression logic."


def _plain_language_for_property(label: str) -> str:
    if "lookup condition" in label.lower():
        return "A record matches the lookup only when this condition is true."
    if "filter condition" in label.lower() or "source filter" in label.lower():
        return "Include/select records only when this condition is true."
    if "join condition" in label.lower() or "user defined join" in label.lower():
        return "Rows are matched only when this join condition is true."
    if "update strategy" in label.lower():
        return "Row operation is assigned according to this update strategy expression."
    return "Business logic rule extracted from transformation properties."


def _eligibility_rows(repo) -> List[Dict[str, object]]:
    rows: List[Dict[str, object]] = []
    prop_labels = [
        "Lookup Condition",
        "Filter Condition",
        "Source Filter",
        "Join Condition",
        "User Defined Join",
        "Update Strategy Expression",
    ]

    for mapping_name, session_name in _ordered_mapping_sessions(repo):
        mapping = repo.mappings.get(mapping_name)
        if mapping is None:
            continue
        for inst in mapping.instances:
            if inst.type != "TRANSFORMATION":
                continue
            tobj = repo.transformation(mapping_name, inst.ref_name)
            if tobj is None:
                continue
            ttype = tobj.type or inst.ref_type or ""

            for expr_info in tobj.expressions:
                expr = (expr_info.get("expression") or "").strip()
                if not _is_rule_expression(expr):
                    continue
                port = expr_info.get("port") or "(port)"
                rows.append({
                    "Session": session_name or "",
                    "Mapping/Mapplet": mapping_name,
                    "Transformation Name": inst.name,
                    "Transformation Type": ttype,
                    "Eligibility Rule/Logic (Technical)": expr,
                    "Eligibility Rule/Logic (Plain Language)": _plain_language_for_expression(str(port)),
                    "Source (Excel/XML)": "XML/JSON (uploaded Informatica export)",
                })

            attrs = tobj.attributes or {}
            for label in prop_labels:
                value = (attrs.get(label) or "").strip()
                if not value:
                    continue
                rows.append({
                    "Session": session_name or "",
                    "Mapping/Mapplet": mapping_name,
                    "Transformation Name": inst.name,
                    "Transformation Type": ttype,
                    "Eligibility Rule/Logic (Technical)": value,
                    "Eligibility Rule/Logic (Plain Language)": _plain_language_for_property(label),
                    "Source (Excel/XML)": "XML/JSON (uploaded Informatica export)",
                })
    return rows


def _eligibility_summary_rows(eligibility_rows: List[Dict[str, object]]) -> List[Dict[str, object]]:
    grouped: Dict[Tuple[str, str], List[str]] = defaultdict(list)
    for row in eligibility_rows:
        key = (str(row.get("Session", "")), str(row.get("Mapping/Mapplet", "")))
        plain = str(row.get("Eligibility Rule/Logic (Plain Language)", "")).strip()
        tech = str(row.get("Eligibility Rule/Logic (Technical)", "")).strip()
        if plain or tech:
            grouped[key].append(f"- {plain}\n  Technical: {tech}")

    out = []
    for (session, mapping), lines in grouped.items():
        out.append({
            "Session": session,
            "Mapping/Mapplet": mapping,
            "Eligibility Rules/Logics": "\n\n".join(lines),
        })
    return out


def generate_informatica_xml_downloads_excel(repo, out_path: str) -> str:
    wb = Workbook()

    ws1 = wb.active
    ws1.title = "Field_Lineage"
    field_rows = _field_lineage_rows(repo)
    _write_sheet(
        ws1,
        [
            "Mapping", "Session", "Source Table", "Source Field", "Target Table", "Target Field",
            "Transformation", "Transformation Type", "Transformation_Full_Lineage_path",
            "Individual_Transformations", "Links", "Hop count",
        ],
        field_rows,
        [28, 26, 28, 24, 28, 24, 28, 22, 90, 54, 36, 10],
    )

    ws2 = wb.create_sheet("Referrential-other Tables")
    ref_rows = _reference_other_tables_rows(repo)
    _write_sheet(
        ws2,
        ["Mapping", "Transformation", "Referrential Table Name", "Alias /Original Name", "Referrential Field"],
        ref_rows,
        [30, 34, 34, 28, 30],
    )

    ws3 = wb.create_sheet("Eligibility Rules")
    eligibility_rows = _eligibility_rows(repo)
    _write_sheet(
        ws3,
        [
            "Session",
            "Mapping/Mapplet",
            "Transformation Name",
            "Transformation Type",
            "Eligibility Rule/Logic (Technical)",
            "Eligibility Rule/Logic (Plain Language)",
            "Source (Excel/XML)",
        ],
        eligibility_rows,
        [26, 30, 32, 22, 92, 64, 44],
    )

    ws4 = wb.create_sheet("Eligibility Rules - Summary")
    summary_rows = _eligibility_summary_rows(eligibility_rows)
    _write_sheet(
        ws4,
        ["Session", "Mapping/Mapplet", "Eligibility Rules/Logics"],
        summary_rows,
        [26, 30, 140],
    )

    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    wb.save(out_path)
    return out_path
