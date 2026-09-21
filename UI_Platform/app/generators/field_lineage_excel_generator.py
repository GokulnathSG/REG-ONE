"""Builds the 'Field_Lineage' Excel workbook.

Sheet: Field_Lineage -- one row per single (Source Field -> Target Field)
dependency (not one row per physical hop), covering every Mapping/Mapplet/
Transformation field in the workflow.

Columns: Session, Mapping, Mapplet, Source Table, Source Field,
         Target Table, Target Field, Transformation, Transformation Type,
         Transformation_Full_Lineage_Path, Individual_Transformations,
         Dependency Type, Links, Hop Count
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

NAVY = "1F3864"
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
HEADER_FILL = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
BODY_FONT = Font(size=10)
WRAP_TOP = Alignment(vertical="top", wrap_text=True)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
STRIPE_FILL = PatternFill(start_color="F2F6FA", end_color="F2F6FA", fill_type="solid")
DIRECT_FILL = PatternFill(start_color="E2F0D9", end_color="E2F0D9", fill_type="solid")
INDIRECT_FILL = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")

COLUMNS = [
    "Session", "Mapping", "Mapplet", "Source Table", "Source Field",
    "Target Table", "Target Field", "Transformation", "Transformation Type",
    "Transformation_Full_Lineage_Path", "Individual_Transformations",
    "Dependency Type", "Links", "Hop Count",
]
COL_WIDTHS = [22, 22, 20, 20, 18, 20, 18, 24, 20, 60, 44, 14, 44, 10]


def generate_field_lineage_excel(rows, out_path: str) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Field_Lineage"

    ws.append(COLUMNS)
    for col_idx in range(1, len(COLUMNS) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)
        cell.border = BORDER

    for r_idx, row in enumerate(rows, start=2):
        dep_type = row.get("Dependency Type", "")
        for c_idx, col_name in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=row.get(col_name, ""))
            cell.font = BODY_FONT
            cell.alignment = WRAP_TOP
            cell.border = BORDER
            if col_name == "Dependency Type" and dep_type == "Direct":
                cell.fill = DIRECT_FILL
            elif col_name == "Dependency Type" and dep_type == "Indirect":
                cell.fill = INDIRECT_FILL
            elif r_idx % 2 == 0:
                cell.fill = STRIPE_FILL

    for c_idx, width in enumerate(COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(c_idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(len(rows) + 1, 1)}"

    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    wb.save(out_path)
    return out_path
