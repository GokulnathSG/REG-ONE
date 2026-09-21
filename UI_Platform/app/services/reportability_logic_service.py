"""Workbook-backed business-logic lookup for the Reportability screen.

The Reportability flow is already single-process and in-memory, so this keeps
the most recently uploaded workbook in a tiny cache keyed by the exact fields
the UI needs when a lineage node is clicked. The Transformations sheet is
required; Mapplet_Transformations is optional.
"""
from __future__ import annotations

from io import BytesIO
import re
from typing import Dict, Optional, Tuple

from openpyxl import load_workbook
from werkzeug.utils import secure_filename

from app.services.workflow_service import UploadError


_CACHE = {"filename": None, "data": None}

_TRANSFORMATION_SHEET = "Transformations"
_MAPPLET_SHEET = "Mapplet_Transformations"

_TRANSFORMATION_HEADERS = {
    "Transformation Name": "transformation_name",
    "PORT_NAME": "port_name",
    "Transformation Type": "transformation_type",
    "Business Logic": "business_logic",
}

_MAPPLET_HEADERS = {
    "Mapplet Name": "mapplet_name",
    "Transformation Name": "transformation_name",
    "PORT_NAME": "port_name",
    "Transformation Type": "transformation_type",
    "Business Logic": "business_logic",
}

_VARIABLE_REF_RE = re.compile(r"^(v_[A-Za-z0-9_$.#@]*)\b", re.IGNORECASE)


def _normalize(value: object) -> str:
    return str(value or "").strip()


def _key_part(value: object) -> str:
    return _normalize(value).casefold()


def _merge_logic(existing: Optional[str], incoming: str) -> str:
    incoming = _normalize(incoming)
    if not incoming:
        return existing or ""
    if not existing:
        return incoming
    seen = {part.strip() for part in existing.split("\n\n") if part.strip()}
    if incoming in seen:
        return existing
    return existing + "\n\n" + incoming


def _header_map(sheet, expected_headers: Dict[str, str]) -> Dict[str, int]:
    header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if header_row is None:
        raise UploadError("INVALID_EXCEL", f"Sheet '{sheet.title}' is empty.")

    positions = {_normalize(value): idx for idx, value in enumerate(header_row)}
    missing = [name for name in expected_headers if name not in positions]
    if missing:
        raise UploadError(
            "INVALID_EXCEL",
            f"Sheet '{sheet.title}' is missing required column(s): {', '.join(missing)}.",
        )
    return {target: positions[source] for source, target in expected_headers.items()}


def _parse_sheet(sheet, expected_headers: Dict[str, str]) -> Tuple[Dict[tuple, str], int]:
    header_map = _header_map(sheet, expected_headers)
    rows = 0
    parsed: Dict[tuple, str] = {}
    for values in sheet.iter_rows(min_row=2, values_only=True):
        data = {name: _normalize(values[idx] if idx < len(values) else "") for name, idx in header_map.items()}
        if not any(data.values()):
            continue
        rows += 1
        logic = data.pop("business_logic", "")
        key = tuple(_key_part(data[name]) for name in data)
        if not any(key):
            continue
        parsed[key] = _merge_logic(parsed.get(key), logic)
    return parsed, rows


def _is_expression_type(transformation_type: str) -> bool:
    return _key_part(transformation_type).startswith("expression")


def _leading_variable_reference(logic: str) -> str:
    match = _VARIABLE_REF_RE.match(_normalize(logic))
    return match.group(1) if match else ""


def _resolve_expression_logic(table: Dict[tuple, str], key: tuple, port_index: int,
                              type_index: int, visited: Optional[set] = None) -> Optional[str]:
    if visited is None:
        visited = set()
    if key in visited:
        return table.get(key)
    visited.add(key)

    logic = table.get(key)
    if not logic or not _is_expression_type(key[type_index]):
        return logic

    variable_name = _leading_variable_reference(logic)
    if not variable_name:
        return logic

    ref_key = list(key)
    ref_key[port_index] = _key_part(variable_name)
    ref_key = tuple(ref_key)
    if ref_key == key:
        return logic

    resolved = _resolve_expression_logic(table, ref_key, port_index, type_index, visited)
    return resolved or logic


def clear_workbook() -> None:
    _CACHE["filename"] = None
    _CACHE["data"] = None


def load_workbook_from_upload(file_storage) -> dict:
    filename = secure_filename(file_storage.filename or "")
    if not filename:
        raise UploadError("INVALID_EXCEL", "Please choose an Excel workbook to upload.")
    lowered = filename.lower()
    if not (lowered.endswith(".xlsx") or lowered.endswith(".xlsm")):
        raise UploadError("INVALID_EXCEL", "Only .xlsx and .xlsm workbooks are supported for Reportability.")

    payload = file_storage.read()
    file_storage.stream.seek(0)
    try:
        workbook = load_workbook(filename=BytesIO(payload), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001
        raise UploadError("INVALID_EXCEL", f"Could not read the Excel workbook: {exc}")

    if _TRANSFORMATION_SHEET not in workbook.sheetnames:
        raise UploadError(
            "INVALID_EXCEL",
            f"Workbook is missing required sheet '{_TRANSFORMATION_SHEET}'.",
        )

    transformations, transformation_rows = _parse_sheet(
        workbook[_TRANSFORMATION_SHEET],
        _TRANSFORMATION_HEADERS,
    )
    if _MAPPLET_SHEET in workbook.sheetnames:
        mapplets, mapplet_rows = _parse_sheet(
            workbook[_MAPPLET_SHEET],
            _MAPPLET_HEADERS,
        )
    else:
        mapplets, mapplet_rows = {}, 0

    data = {
        "transformations": transformations,
        "mapplets": mapplets,
        "transformation_rows": transformation_rows,
        "mapplet_rows": mapplet_rows,
    }
    _CACHE["filename"] = filename
    _CACHE["data"] = data
    return data


def get_workbook_name() -> Optional[str]:
    return _CACHE["filename"]


def get_transformation_logic(transformation_name: str, port_name: str, transformation_type: str) -> Optional[str]:
    data = _CACHE.get("data") or {}
    key = (_key_part(transformation_name), _key_part(port_name), _key_part(transformation_type))
    table = data.get("transformations") or {}
    return _resolve_expression_logic(table, key, port_index=1, type_index=2)


def get_transformation_expression_reference(transformation_name: str, port_name: str,
                                            transformation_type: str) -> Optional[str]:
    data = _CACHE.get("data") or {}
    key = (_key_part(transformation_name), _key_part(port_name), _key_part(transformation_type))
    table = data.get("transformations") or {}
    logic = table.get(key)
    if not logic:
        return None
    return _leading_variable_reference(logic)


def get_mapplet_logic(mapplet_name: str, transformation_name: str, port_name: str,
                      transformation_type: str) -> Optional[str]:
    data = _CACHE.get("data") or {}
    key = (
        _key_part(mapplet_name),
        _key_part(transformation_name),
        _key_part(port_name),
        _key_part(transformation_type),
    )
    table = data.get("mapplets") or {}
    return _resolve_expression_logic(table, key, port_index=2, type_index=3)


def get_mapplet_expression_reference(mapplet_name: str, transformation_name: str, port_name: str,
                                     transformation_type: str) -> Optional[str]:
    data = _CACHE.get("data") or {}
    key = (
        _key_part(mapplet_name),
        _key_part(transformation_name),
        _key_part(port_name),
        _key_part(transformation_type),
    )
    table = data.get("mapplets") or {}
    logic = table.get(key)
    if not logic:
        return None
    return _leading_variable_reference(logic)